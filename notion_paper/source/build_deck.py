"""Build an editable research deck and speaker handout from the reviewed content JSON.

Run with the existing MetaFind Python environment. No model/data producers run here.
"""
from pathlib import Path
import hashlib,json,math,re
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
CONTENT=Path(__file__).with_name('ppt_content.json')
ASSETS=ROOT/'assets';ASSETS.mkdir(exist_ok=True)
FONT='Noto Sans CJK TC'
INK='183247';TEAL='137D88';ORANGE='CE663F';MUTED='657A88';BG='F5F7F8';PALE='E6F1F0';LINE='D9E2E7';WHITE='FFFFFF'
COLORS=[INK,TEAL,ORANGE,'8E9FAA','8778A9']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='MetaFind｜標註更新與評估計畫'
prs.core_properties.subject='最新標註方法、完成時間與後續評估計畫；含逐頁講稿'
prs.core_properties.author='MetaFindV1'
prs.core_properties.keywords='MetaFind, Table 1, Stage 2, reproduction, 自訂評估'
shape_boxes=[]

def clean(v):
    return re.sub(r'`([^`]+)`',r'\1',str(v)).replace('**','')

def setfont(run,size,color=INK,bold=False):
    run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    rp=run._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'):
        node=rp.find(tag, rp.nsmap)
        if node is None:node=OxmlElement(tag);rp.append(node)
        node.set('typeface',FONT)

def text(slide,value,x,y,w,h,size=22,color=INK,bold=False,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.015);tf.margin_right=Inches(.015)
    tf.margin_top=Inches(.01);tf.margin_bottom=Inches(.01)
    for i,line in enumerate(clean(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
        p.space_before=Pt(0);p.space_after=Pt(2);p.line_spacing=1.12
        r=p.add_run();r.text=line;setfont(r,size,color,bold)
    shape_boxes.append({'slide':len(prs.slides),'text':clean(value),'box':[x,y,w,h],'size':size})
    return shape

def box(slide,x,y,w,h,fill=WHITE,line=None,rounded=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(.7)
    else:sh.line.fill.background()
    if rounded:sh.adjustments[0]=.08
    return sh

def line(slide,x1,y1,x2,y2,color=LINE,width=1.3):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
    return sh

def short_source(s):
    n=int(s['id'])
    if n==5:return '進度快照：2026-09-08 04:12:30（台北）；ETA 為推估｜完整依據見備忘稿'
    if n<6:return '來源：目前標註程式、實際資料與論文 Figure 2｜完整依據見備忘稿'
    if n==6:return '來源：ULIP-2 實際 NPY；插圖為本地同 UID 渲染｜完整欄位與來源見備忘稿'
    return '來源：論文 Table 1 與本地同資產評估規則；正式評估尚待執行｜完整依據見備忘稿'

def base(s,n,total):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    box(sl,0,0,.15,7.5,TEAL,rounded=False)
    text(sl,str(s.get('section','RESEARCH BRIEF')).upper(),.57,.24,11.4,.27,11,TEAL,True)
    title=clean(s['title']);size=31 if len(title)<=30 else 27
    text(sl,title,.57,.69,12.16,1.01,size,INK,True)
    line(sl,.6,1.62,12.73,1.62)
    box(sl,.6,6.47,12.1,.42,PALE,rounded=False)
    take=clean(s['takeaway']);text(sl,take,.75,6.50,11.8,.34,15 if len(take)>48 else 17,TEAL,True)
    footer=text(sl,short_source(s),.60,7.08,11.1,.20,8.2,MUTED)
    text(sl,f'{n:02d} / {total:02d}',11.78,7.05,.93,.26,11,MUTED,align=PP_ALIGN.RIGHT)
    return sl

def bullets(sl,items,x,y,w,step=.95,size=22):
    for i,item in enumerate(items):
        text(sl,'•',x,y+i*step,.2,.4,size,TEAL,True)
        text(sl,item,x+.3,y+i*step,w-.3,step-.06,size,INK)

def cards(sl,s):
    items=s.get('bullets',[])
    cols=2 if len(items)>2 else max(len(items),1)
    rows=math.ceil(max(len(items),1)/cols)
    w=(12.1-.28*(cols-1))/cols;h=(4.48-.25*(rows-1))/rows
    for i,item in enumerate(items):
        x=.6+(i%cols)*(w+.28);y=1.82+(i//cols)*(h+.25)
        box(sl,x,y,w,h,WHITE,LINE)
        text(sl,f'{i+1:02d}',x+.23,y+.18,.73,.42,23,TEAL,True)
        text(sl,item,x+.24,y+.82,w-.48,h-.97,23 if len(clean(item))<45 else 20,INK)

def draw_table(sl,s):
    data=s['table'];headers=data['headers'];rows=data['rows'];count=len(headers)
    ratios=data.get('widths')
    if not ratios or len(ratios)!=count:
        ratios=([2.5]+[1]*(count-1)) if count>=6 else ([1.25]+[1]*(count-1))
    widths=[12.1*float(v)/sum(ratios) for v in ratios]
    top=s.get('_table_top',1.85);available=s.get('_table_height',3.67 if s.get('bullets') else 4.4)
    heights=[.54]+[(available-.54)/max(len(rows),1)]*len(rows)
    y=top
    for ri,row in enumerate([headers]+rows):
        x=.6
        assert len(row)==count,(s['id'],ri,row)
        for ci,value in enumerate(row):
            fill=INK if ri==0 else (WHITE if ri%2 else 'EBF0F3')
            box(sl,x,y,widths[ci],heights[ri],fill,rounded=False)
            raw=clean(value);size=s.get('_table_size',14 if count>=6 else 17)
            if len(raw)>60:size=14
            if len(raw)>90:size=12.5
            shape=text(sl,raw,x+.10,y+.06,widths[ci]-.2,heights[ri]-.1,size,WHITE if ri==0 else INK,ri==0 or ci==0)
            shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
            x+=widths[ci]
        y+=heights[ri]
    if s.get('bullets') and not s.get('_no_table_bullets'):bullets(sl,s['bullets'][:2],.7,5.70,11.9,.31,14)


def picture(sl,name,x,y,w,h):
    path=ASSETS/name
    with Image.open(path) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);sw,sh=iw*scale,ih*scale
    sl.shapes.add_picture(str(path),Inches(x+(w-sw)/2),Inches(y+(h-sh)/2),width=Inches(sw),height=Inches(sh))

def table(sl,data,top=1.94,height=4.17,size=18,widths=None):
    data={**data}
    if widths:data['widths']=widths
    draw_table(sl,{'id':'table','table':data,'_table_top':top,'_table_height':height,'_table_size':size})

def card(sl,number,title,detail,x,y,w,h,dark=False):
    box(sl,x,y,w,h,INK if dark else WHITE,None if dark else LINE)
    text(sl,number,x+.2,y+.14,.7,.35,16,'8ED3CF' if dark else TEAL,True)
    text(sl,title,x+.2,y+.61,w-.4,.76,22,WHITE if dark else INK,True)
    text(sl,detail,x+.2,y+1.47,w-.4,h-1.58,17,'C6D8E1' if dark else MUTED)

def cover(sl,s):
    sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(INK)
    text(sl,'METAFIND  /  RESEARCH UPDATE',.73,.62,11.8,.4,13,'8ED3CF',True)
    text(sl,'標註更新\n與評估計畫',.72,1.68,7.05,1.97,45,WHITE,True)
    text(sl,'現在怎麼做，完成後怎麼驗證',.77,4.04,7.5,.65,24,'C6D8E1')
    text(sl,'2026.09.08  ·  10 頁  ·  約 12–13 分鐘',.77,6.76,10.6,.3,12,'A9C2CD')
    for i,(title,detail) in enumerate([('目前的標註','方法・資料卡・完成時間'),('接下來的評估','測試設計・差異・後續流程')]):
        y=1.72+i*2.18;box(sl,8.66,y,3.91,1.83,'23495C')
        text(sl,f'0{i+1}',8.89,y+.18,.55,.33,15,'8ED3CF',True)
        text(sl,title,8.9,y+.67,3.42,.50,24,WHITE,True)
        text(sl,detail,8.9,y+1.28,3.42,.29,12,'C6D8E1')

def render(sl,s,n):
    if n==2:
        for i,st in enumerate(s['steps']):
            x=.6+i*3.095;box(sl,x,1.89,2.82,1.58,WHITE,LINE)
            text(sl,f'0{i+1}',x+.17,2.04,.45,.3,15,TEAL,True)
            text(sl,st['title'],x+.17,2.48,2.48,.43,20,INK,True)
            text(sl,st['detail'],x+.17,3.01,2.48,.32,12,MUTED)
            if i<3:text(sl,'→',x+2.85,2.50,.24,.3,17,TEAL)
        for i in range(3):picture(sl,f'annotation_view_{i+1:02d}.png',.6+i*2.47,3.77,2.23,2.23)
        text(sl,'同一張椅子\n從不同角度看',8.36,4.10,3.9,1.12,27,INK,True)
        text(sl,'示意選 3 張；實際輸入 11 張',8.36,5.50,3.95,.39,16,TEAL)
    elif n==3:
        table(sl,s['table'],top=1.93,height=3.58,size=18,widths=[1.05,1.6,2.35])
        text(sl,'完整資料卡共 13 欄；尺寸、重量是模型估計，並非量測真值。',.76,5.90,11.65,.36,18,ORANGE,True)
    elif n==4:
        stages=[('檢查資料卡','缺欄、格式、尺寸與重量範圍'),('帶著錯誤重答','初次＋最多兩次修正'),('收錄或隔離','合格收錄；失敗保留原因')]
        for i,(title,detail) in enumerate(stages):card(sl,f'0{i+1}',title,detail,.6+i*4.12,1.95,3.86,2.35)
        table(sl,{'headers':['項目','論文','目前做法'],'rows':[['標註模型','GPT-4o','Gemma 4 12B・BF16'],['多視角輸入','11 張','11 張；提示與重試為本地選擇']]},top=4.62,height=1.53,size=16,widths=[1.1,1.2,2.5])
    elif n==5:
        p=json.loads((ROOT/'source/annotation_progress_snapshot.json').read_text());c=p['counts']
        text(sl,f"{c['processed_percent_of_queue']:.2f}%",.67,1.93,5.45,.94,53,TEAL,True)
        text(sl,f"本輪已處理 {c['processed_this_run_success_plus_failure']:,} / {c['scheduled_this_run']:,}",.72,3.04,7.0,.48,23,INK,True)
        box(sl,.72,3.82,7.20,.26,LINE,rounded=False)
        box(sl,.72,3.82,7.20*c['processed_percent_of_queue']/100,.26,TEAL,rounded=False)
        for i,(label,value) in enumerate([('成功',c['successfully_written_this_run']),('失敗隔離',c['failed_this_run_unique_uids']),('待處理',c['remaining_scheduled'])]):
            x=.72+i*2.52;text(sl,label,x,4.43,2.2,.35,16,MUTED);text(sl,f'{value:,}',x,5.00,2.2,.64,29,INK,True)
        box(sl,8.55,1.96,4.14,4.10,INK)
        text(sl,'預計完成本輪標註',8.83,2.27,3.57,.41,18,'8ED3CF',True)
        text(sl,'9/9 上午',8.82,3.10,3.64,.85,34,WHITE,True)
        text(sl,'規劃 09:00–12:00',8.84,4.14,3.61,.43,21,WHITE)
        text(sl,'近期約 11.7–11.9 件／分鐘\n需持續運作且速度相近',8.84,5.00,3.55,.79,15,'C6D8E1')
        text(sl,'只估標註處理；不含後續編碼、訓練與評估。',.76,6.12,11.65,.26,13,MUTED)
    elif n==6:
        raw=json.loads((ROOT/s['example_source']).read_text())['record']
        box(sl,.6,1.89,3.05,4.18,WHITE,LINE)
        text(sl,raw['text'][0],.79,2.05,2.65,.43,23,INK,True)
        text(sl,'Objaverse · 15 個欄位',.79,2.58,2.65,.29,12,TEAL)
        picture(sl,'annotation_view_01.png',.82,2.98,2.60,2.34)
        text(sl,'同 UID 的本地渲染圖',.79,5.45,2.65,.29,12,MUTED)
        text(sl,'UID：85059770…',.79,5.82,2.65,.23,11,MUTED)
        for i,(label,field,translation) in enumerate([
            ('BLIP 原始描述','blip_caption','粉紅坐墊、木框的椅子'),
            ('MSFT 原始描述','msft_caption','紅色坐墊的椅子')]):
            y=1.89+i*1.26;box(sl,3.91,y,8.79,1.13,WHITE,LINE)
            text(sl,label,4.11,y+.08,8.34,.27,12,TEAL,True)
            text(sl,raw[field],4.11,y+.40,8.34,.37,19,INK,True)
            text(sl,translation,4.11,y+.84,8.34,.25,13,MUTED)
        for i,(label,detail) in enumerate([
            ('原始點資料','xyz / rgb\n各 10000 × 3'),
            ('圖片特徵','image_feat：12 × 1280\nthumbnail_feat：1280'),
            ('文字與特徵','名稱／BLIP／MSFT 向量\n另有 16 段 retrieval_text')]):
            x=3.91+i*2.97;box(sl,x,4.54,2.85,1.53,PALE)
            text(sl,label,x+.14,4.70,2.57,.30,16,TEAL,True)
            text(sl,detail,x+.14,5.18,2.57,.71,13.5,INK)
        text(sl,'這筆沒有原始圖片、GLB、尺寸或重量；完整 15 欄保留於來源檔。',.74,6.14,11.9,.26,13,ORANGE)
    elif n==7:
        for i,st in enumerate(s['steps']):
            card(sl,f'0{i+1}',st['title'],st['detail'],.6+i*4.12,1.94,3.86,2.89,dark=i==2)
        text(sl,'比較模型',.75,5.12,1.38,.37,17,TEAL,True)
        text(sl,'ULIP-2＋平均融合　／　Stage 1　／　Stage 2-off',2.22,5.09,10.35,.44,21,INK,True)
        text(sl,'七種條件',.75,5.86,1.38,.33,17,TEAL,True)
        text(sl,'T　I　P　T＋I　T＋P　I＋P　全部',2.22,5.85,10.35,.36,20,INK)
    elif n==8:
        example=s['metric_example'];count=example['queries']
        text(sl,f'教學示意：{count} 道 query，每題對應 1 個目標 UID',.76,1.97,11.81,.46,23,TEAL,True)
        for i,(name,hits,detail) in enumerate([
            ('R@1',example['top1_successes'],'第一名就是目標物件'),
            ('R@5',example['top5_successes'],'前五名包含目標物件')]):
            x=.6+i*6.19;box(sl,x,2.72,5.91,2.50,WHITE,LINE)
            text(sl,name,x+.23,2.94,5.45,.48,27,TEAL,True)
            text(sl,f'{100*hits/count:g}%',x+.22,3.55,2.44,.80,45,INK,True)
            text(sl,f'{hits} ÷ {count}',x+3.02,3.77,2.43,.39,24,MUTED)
            text(sl,detail,x+.23,4.63,5.45,.35,20,INK)
        text(sl,'R@5 的成功題目，包含已在第一名找對的題目。',.77,5.56,11.75,.36,19,INK)
        text(sl,'30%／65% 僅用來說明計算，並非模型實驗結果。',.77,6.08,11.75,.27,14,ORANGE)
    elif n==9:
        table(sl,s['table'],top=1.90,height=4.32,size=16,widths=[1.04,1.68,2.60])
    elif n==10:
        for i,st in enumerate(s['steps']):
            x,y=.6+(i%3)*4.12,1.91+(i//3)*2.21
            box(sl,x,y,3.86,1.98,INK if i==5 else WHITE,None if i==5 else LINE)
            text(sl,f'0{i+1}',x+.16,y+.12,.48,.30,15,'8ED3CF' if i==5 else TEAL,True)
            text(sl,st['title'],x+.16,y+.57,3.54,.44,20,WHITE if i==5 else INK,True)
            text(sl,st['detail'],x+.16,y+1.14,3.54,.67,15,'C6D8E1' if i==5 else MUTED)
            if i%3<2:text(sl,'→',x+3.88,y+.82,.23,.32,17,TEAL)

def make():
    obj=json.loads(CONTENT.read_text());slides=obj['slides'];assert len(slides)==10
    manuscript=['# '+obj['meta']['title']+' — 逐頁講稿','','10 頁；約 12–13 分鐘。以下講稿逐頁寫入 PPT 備忘稿。最新完整資料的正式評估尚未執行。','']
    for n,s in enumerate(slides,1):
        if n==1:sl=prs.slides.add_slide(prs.slide_layouts[6]);cover(sl,s)
        else:sl=base(s,n,10);render(sl,s,n)
        seconds=s.get('duration_seconds',75);sources='\n'.join(s['sources'])
        notes=f"第 {n:02d} 頁｜{s['title']}\n建議時間：{seconds} 秒\n\n{s['notes']}\n\n本頁重點：{s['takeaway']}\n\n來源與界線\n{sources}"
        sl.notes_slide.notes_text_frame.text=notes
        for p in sl.notes_slide.notes_text_frame.paragraphs:
            for r in p.runs:setfont(r,12)
        manuscript += [f"## {n:02d}｜{s['title']}",'',f'**建議時間：{seconds} 秒。**','',s['notes'],'','**本頁重點：** '+s['takeaway'],'','來源：','']
        manuscript += ['- '+v for v in s['sources']]+['']
    out=ROOT/'MetaFind_Table1_Stage2_含講稿.pptx';prs.save(out)
    (ROOT/'逐頁講稿.md').write_text('\n'.join(manuscript).rstrip()+'\n')
    (ROOT/'source/layout_boxes.json').write_text(json.dumps(shape_boxes,ensure_ascii=False,indent=2)+'\n')
    record={'slides':10,'notes_slides':10,'content_sha256':hashlib.sha256(CONTENT.read_bytes()).hexdigest(),'pptx_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'font':FONT,'size':'16:9','environment':'existing MetaFind Python; python-pptx 1.0.2; no model/data producers','images':'three unchanged actual renders of one annotated chair; source/source_manifest.json'}
    (ROOT/'source/build_record.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'pptx':str(out),'slides':10,'bytes':out.stat().st_size},ensure_ascii=False))

if __name__=='__main__':make()
